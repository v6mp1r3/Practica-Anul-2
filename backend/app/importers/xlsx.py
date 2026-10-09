"""The same weekly grid as an Excel workbook (groups across the top, days and pair times down the left).
Merged cells say what the drawn borders say in a PDF: a cell merged across several groups is a lecture they
share, a pair's cell merged over its rows is every week, and a pair split into an upper and a lower cell holds
the odd week above and the even week below."""
import io
import re

from .cells import TIME_RE, day_of, group_name, read_cell, untyped_are_seminars


def _time(v) -> tuple[str, str] | None:
    s = str(v or "")
    m = TIME_RE.search(s) or re.search(r"(\d{1,2})(\d{2})\s*[-–]\s*(\d{1,2})(\d{2})", s)
    return (f"{int(m[1]):02d}:{m[2]}", f"{int(m[3]):02d}:{m[4]}") if m else None


def read_xlsx(data: bytes) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    rows: list[dict] = []
    warnings: list[str] = []
    for ws in wb.worksheets:
        # each cell → the top-left cell of its merged range, and that range
        anchor: dict[tuple[int, int], tuple[int, int, int, int]] = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    anchor[(r, c)] = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)

        def box(r: int, c: int) -> tuple[int, int, int, int]:
            return anchor.get((r, c), (r, c, r, c))

        def value(r: int, c: int):
            a = box(r, c)
            return ws.cell(a[0], a[1]).value

        # the header: the row with the most group names
        best: list[tuple[int, str]] = []
        header_row = 0
        for r in range(1, min(ws.max_row, 40) + 1):
            found = [(c, n) for c in range(1, ws.max_column + 1) if (n := group_name(str(ws.cell(r, c).value or "")))]
            if len(found) > len(best):
                best, header_row = found, r
        if not best:
            continue
        first_col = min(c for c, _ in best)
        # the group columns: a group's header may be merged over several columns
        col_group: dict[int, str] = {}
        for c, n in best:
            _, c0, _, c1 = box(header_row, c)
            for x in range(c0, c1 + 1):
                col_group[x] = n
        # pairs down the left: a time cell (merged over the pair's rows), under the last day name above it
        bands = []
        day = None
        r = header_row + 1
        while r <= ws.max_row:
            for c in range(1, first_col):
                d = day_of(str(value(r, c) or ""))
                if d is not None:
                    day = d
            t = next((tm for c in range(1, first_col) if (tm := _time(value(r, c)))), None)
            if t and day is not None:
                c = next(c for c in range(1, first_col) if _time(value(r, c)))
                r0, _, r1, _ = box(r, c)
                bands.append((r0, r1, t[0], t[1], day))
                r = r1 + 1
                continue
            r += 1
        if not bands:
            warnings.append(f"{ws.title}: no pair times found")
            continue
        for r0, r1, start, end, day in bands:
            seen: set[tuple[int, int]] = set()
            for c in sorted(col_group):
                cells = []
                for r in range(r0, r1 + 1):
                    a = box(r, c)
                    if (a[0], a[1]) in seen or ws.cell(a[0], a[1]).value in (None, ""):
                        continue
                    seen.add((a[0], a[1]))
                    cells.append(a)
                if not cells:
                    continue
                # the groups under the widest of these cells (a lecture merged across groups)
                c0 = min(a[1] for a in cells)
                c1 = max(a[3] for a in cells)
                groups = list(dict.fromkeys(col_group[x] for x in range(c0, c1 + 1) if x in col_group))
                mid = r0 + (r1 - r0 + 1) / 2
                top = [a for a in cells if a[2] < mid]
                bottom = [a for a in cells if a[0] >= mid]
                text = lambda cs: [line for a in sorted(cs) for line in str(ws.cell(a[0], a[1]).value).splitlines() if line.strip()]
                parts = [(cells, "weekly")]
                if r1 > r0 and len(top) + len(bottom) == len(cells):
                    full = lambda cs: [x for x in read_cell(text(cs)) if x["teacher"] or x["room"]]
                    if top and bottom and full(top) and full(bottom):
                        parts = [(top, "odd"), (bottom, "even")]
                    elif r1 - r0 == 1 and (top or bottom):
                        parts = [(top, "odd")] if top else [(bottom, "even")]
                for cs, parity in parts:
                    for x in read_cell(text(cs)):
                        rows.append({"groups": groups, "day": day, "start": start, "end": end, "parity": parity, **x,
                                     "source": f"{ws.title} {ws.cell(cs[0][0], cs[0][1]).coordinate}"})
    return {"rows": untyped_are_seminars(rows), "warnings": warnings}
