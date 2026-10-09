"""Reading the faculties' own timetables (PDF and Excel) into rows. The PDF is FIMIT's real first-year
timetable for the 2025-2026 spring semester, as published on fimit.utm.md."""
import io
from pathlib import Path

import pytest

from app.importers import read_sheet
from app.importers.cells import clean_teacher, group_name, read_cell

pytestmark = pytest.mark.no_db
PDF = Path(__file__).parent / "fixtures" / "utm_fimit_anul1.pdf"


def _find(rows, day, start, group, subject):
    return [r for r in rows if r["day"] == day and r["start"] == start and group in r["groups"] and subject.lower() in r["subject"].lower()]


def test_cell_one_line_per_field():
    assert read_cell(["Seminar", "Planificarea și Infrastructura T.U.", "asist. univ. Rotaru Igor", "6 - 320A"]) == [
        {"type": "seminar", "subject": "Planificarea și Infrastructura T.U.", "teacher": "Rotaru Igor", "room": "6 - 320A"}
    ]


def test_cell_squeezed_onto_one_line():
    assert read_cell(["VEPS curs E. Guțuleac aud 614"]) == [{"type": "lecture", "subject": "VEPS", "teacher": "E. Guțuleac", "room": "614"}]


def test_cell_type_after_the_subject_and_a_second_teacher():
    [c] = read_cell(["ANALIZA MATEMATICĂ II          Curs", "conf. univ. Rusu Elena", "asist. univ. Mihailov L.", "5-I"])
    assert (c["type"], c["subject"], c["teacher"], c["room"]) == ("lecture", "ANALIZA MATEMATICĂ II", "Rusu Elena", "5-I")


def test_names():
    assert group_name("TDF - 241") == "TDF-241" and group_name("IMCM 241") == "IMCM-241" and group_name("Grupele") is None
    assert clean_teacher("asist.univ. Zubac Vadim") == "Zubac Vadim"
    assert clean_teacher("d.ş.e. Vîrcolici Margareta") == "Vîrcolici Margareta"


def test_utm_pdf_monday():
    rows = read_sheet("orar.pdf", PDF.read_bytes())["rows"]
    # 8:00: Analiza matematică for six groups in the even week only (the upper half of the cell is empty)
    [r] = _find(rows, 0, "08:00", "AR-251", "Analiza")
    assert r["parity"] == "even" and r["groups"] == ["AR-251", "TCM-251", "TDF-251", "IASM-251", "IM-251", "MIF-251"]
    assert (r["type"], r["teacher"], r["end"]) == ("lecture", "Rusu Elena", "09:30")
    # 9:45: the same lecture every week
    assert _find(rows, 0, "09:45", "MIF-251", "Analiza")[0]["parity"] == "weekly"
    # 11:30: a cell split into the odd and the even week
    [odd] = _find(rows, 0, "11:30", "MIF-251", "Analiza")
    [even] = _find(rows, 0, "11:30", "MIF-251", "Teoriei Frigului")
    assert (odd["parity"], odd["type"], odd["room"]) == ("odd", "seminar", "6-309")
    assert even["parity"] == "even"
    # 13:30 and 15:15
    assert _find(rows, 0, "13:30", "IASM-251", "Bazele Agronomiei")[0]["room"] == "15-213"
    assert _find(rows, 0, "15:15", "IMCM-251", "FRANCEZA")[0]["groups"] == ["IM-251", "MIF-251", "IMT-251", "IMCM-251"]
    # Tuesday 8:00: one lecture for all eight groups
    assert len(_find(rows, 1, "08:00", "IMT-251", "DESEN TEHNIC")[0]["groups"]) == 8


def test_xlsx_grid():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["C1"], ws["D1"], ws["E1"] = "Grupele", "TI-231", "TI-232", "FAF-231"
    ws["A3"] = "LUNI"
    ws.merge_cells("A3:A8")
    for i, t in enumerate(["8:00-9:30", "9:45-11:15", "11:30-13:00"]):
        ws.cell(3 + 2 * i, 2, t)
        ws.merge_cells(start_row=3 + 2 * i, start_column=2, end_row=4 + 2 * i, end_column=2)
    ws["C3"] = "Curs\nProgramarea calculatoarelor\nconf. univ. Popescu Ion\n3-101"
    ws.merge_cells("C3:D4")
    ws["C5"] = "Laborator\nStructuri de date\nasist. univ. Ionescu A.\n3-112"
    ws["C6"] = "Laborator\nBaze de date\nasist. univ. Moraru V.\n3-114"
    ws["D7"] = "Seminar\nFizica\nconf. univ. Ceban M.\n1-210"
    buf = io.BytesIO()
    wb.save(buf)
    rows = read_sheet("orar.xlsx", buf.getvalue())["rows"]
    got = [(r["start"], tuple(r["groups"]), r["parity"], r["type"], r["subject"]) for r in rows]
    assert got == [
        ("08:00", ("TI-231", "TI-232"), "weekly", "lecture", "Programarea calculatoarelor"),
        ("09:45", ("TI-231",), "odd", "lab", "Structuri de date"),
        ("09:45", ("TI-231",), "even", "lab", "Baze de date"),
        ("11:30", ("TI-232",), "odd", "seminar", "Fizica"),
    ]


def test_other_files_are_refused():
    from app.errors import ApiError

    with pytest.raises(ApiError):
        read_sheet("orar.docx", b"hello")
